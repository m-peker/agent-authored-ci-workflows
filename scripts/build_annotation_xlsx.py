"""Build the manual-annotation workbook from data/annotation/sample.jsonl (authorship stays hidden).

Sheet "guide": labelling instructions (Turkish). Sheet "items": one row per finding with code snippet,
link and dropdown label columns. Output: data/annotation/annotation_<n>.xlsx
Usage: python scripts/build_annotation_xlsx.py
"""
import json

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

SRC = "data/annotation/sample.jsonl"
items = [json.loads(l) for l in open(SRC, encoding="utf-8")]
OUT = f"data/annotation/annotation_{len(items)}.xlsx"

GUIDE = [
    ("Amaç", "Statik araçların (actionlint, zizmor) bulduğu sorunların gerçekten sorun olup olmadığını doğrulamak."),
    ("Kör etiketleme", "Dosyanın bir yapay zekâ ajanı mı yoksa insan mı tarafından yazıldığı gösterilmiyor. Lütfen GitHub "
                       "linkinde commit/PR yazarına bakmayın; yalnızca dosya içeriğine bakın."),
    ("label", "TP = gerçek sorun | FP = yanlış alarm (sorun değil) | UNSURE = karar verilemiyor"),
    ("breaking için TP", "Workflow çalışmaz, yanlış çalışır ya da GitHub dosyayı reddeder (ör. tanımsız output, hatalı "
                         "needs, YAML hatası, geçersiz izin). Sadece stil/eskimiş sürüm ise FP."),
    ("security_high için TP", "Gerçekçi bir saldırgan bunu kullanabilir (ör. güvenilmeyen PR başlığı run: içinde, "
                              "pull_request_target + PR kodunu checkout, gereksiz write izinleri). Girdi güvenilir "
                              "kaynaktan geliyorsa (ör. github.sha, workflow_dispatch girdisi yalnız bakımcıdan) FP."),
    ("impact (opsiyonel)", "TP ise: HIGH = çalışmayı tamamen bozar / doğrudan istismar edilebilir; "
                           "MEDIUM = belirli koşullarda; LOW = küçük etki"),
    ("note", "Kısa gerekçe (özellikle FP ve UNSURE için faydalı)."),
    ("Süre", "Madde başına ~30-60 sn. Snippet yetmezse GitHub linkini açın; '>>' işaretli satır bulgunun satırıdır."),
]

wb = Workbook()
g = wb.active
g.title = "guide"
g.column_dimensions["A"].width, g.column_dimensions["B"].width = 24, 120
for i, (k, v) in enumerate(GUIDE, 1):
    g.cell(i, 1, k).font = Font(bold=True)
    g.cell(i, 2, v).alignment = Alignment(wrap_text=True, vertical="top")

ws = wb.create_sheet("items")
cols = [("item_id", 8), ("category", 13), ("rule", 20), ("message", 45), ("snippet", 95), ("url", 18),
        ("label", 9), ("impact", 9), ("note", 35)]
for j, (c, w) in enumerate(cols, 1):
    ws.cell(1, j, c).font = Font(bold=True, color="FFFFFF")
    ws.cell(1, j).fill = PatternFill("solid", fgColor="305496")
    ws.column_dimensions[ws.cell(1, j).column_letter].width = w
ws.freeze_panes = "B2"
mono = Font(name="Consolas", size=9)
fill_in = PatternFill("solid", fgColor="FFF2CC")
for i, it in enumerate(items, 2):
    vals = [it["item_id"], it["category"], f'{it["tool"]}:{it["rule"]}', it["message"], it["snippet"], "aç"]
    for j, v in enumerate(vals, 1):
        c = ws.cell(i, j, v)
        c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.cell(i, 5).font = mono
    ws.cell(i, 6).hyperlink = it["url"]
    ws.cell(i, 6).font = Font(color="0563C1", underline="single")
    for j in (7, 8, 9):
        ws.cell(i, j).fill = fill_in
        ws.cell(i, j).alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[i].height = min(409, 13 * (it["snippet"].count("\n") + 1))
n = len(items) + 1
dv = DataValidation(type="list", formula1='"TP,FP,UNSURE"', allow_blank=True)
dv2 = DataValidation(type="list", formula1='"HIGH,MEDIUM,LOW"', allow_blank=True)
ws.add_data_validation(dv)
ws.add_data_validation(dv2)
dv.add(f"G2:G{n}")
dv2.add(f"H2:H{n}")
wb.save(OUT)
print("saved", OUT, len(items), "items")
